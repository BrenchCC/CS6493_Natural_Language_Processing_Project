import json
import logging
import argparse
from pathlib import Path

from pptx import Presentation
from playwright.sync_api import sync_playwright

logger = logging.getLogger(__name__)


DEFAULT_DECK_DIR = Path(__file__).resolve().parent
PPTX_NAME = "CS6493_Math_Reasoning_Presentation.pptx"


def get_expected_slide_count(deck_dir):
    """Read the expected slide count from the slide manifest.

    Args:
        deck_dir: Directory containing slide_manifest.json.

    Returns:
        int: Expected number of slides.
    """
    manifest_path = deck_dir / "slide_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding = "utf-8"))
    return len(manifest)


def parse_args():
    """Parse command-line arguments.

    Returns:
        argparse.Namespace: Parsed command-line options.
    """
    parser = argparse.ArgumentParser(
        description = "Validate the CS6493 HTML deck and editable PPTX."
    )
    parser.add_argument(
        "--deck-dir",
        type = Path,
        default = DEFAULT_DECK_DIR,
        help = "Directory containing index.html and the generated PPTX."
    )
    parser.add_argument(
        "--screenshots-dir",
        type = Path,
        default = DEFAULT_DECK_DIR / "screenshots",
        help = "Directory where validation screenshots should be written."
    )
    return parser.parse_args()


def validate_pptx(pptx_path, expected_slides):
    """Validate the editable PPTX structure.

    Args:
        pptx_path: Path to the generated PowerPoint file.
        expected_slides: Expected slide count from the manifest.

    Returns:
        dict: Summary of slide and shape counts.
    """
    presentation = Presentation(pptx_path)
    pictures = 0
    text_shapes = 0
    charts = 0
    tables = 0
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.shape_type == 13:
                pictures += 1
            if getattr(shape, "has_text_frame", False):
                text_shapes += 1
            if getattr(shape, "has_chart", False):
                charts += 1
            if getattr(shape, "has_table", False):
                tables += 1
    summary = {
        "slides": len(presentation.slides),
        "pictures": pictures,
        "text_shapes": text_shapes,
        "charts": charts,
        "tables": tables,
    }
    if summary["slides"] != expected_slides:
        raise RuntimeError(f"Expected {expected_slides} slides, found {summary['slides']}.")
    if pictures != 0:
        raise RuntimeError(f"Expected no full-slide pictures, found {pictures} picture shapes.")
    return summary


def render_html(deck_path, screenshots_dir, expected_slides):
    """Render representative HTML screenshots and detect obvious overflow.

    Args:
        deck_path: Path to the generated HTML deck.
        screenshots_dir: Directory where screenshots should be written.
        expected_slides: Expected slide count from the manifest.

    Returns:
        dict: Rendering summary with slide count, screenshot names, and overflow details.
    """
    screenshots_dir.mkdir(parents = True, exist_ok = True)
    for old_screenshot in screenshots_dir.glob("slide_*.png"):
        old_screenshot.unlink()
    chrome_path = Path("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome")
    launch_kwargs = {"headless": True}
    if chrome_path.exists():
        launch_kwargs["executable_path"] = str(chrome_path)
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(**launch_kwargs)
        page = browser.new_page(viewport = {"width": 1600, "height": 900}, device_scale_factor = 1)
        page.goto(deck_path.as_uri(), wait_until = "networkidle")
        total = page.locator(".deck > .slide").count()
        overflow = {}
        for index in range(1, total + 1):
            page.goto(deck_path.as_uri() + f"#/{index}", wait_until = "networkidle")
            page.wait_for_timeout(150)
            if index in {1, min(12, expected_slides), expected_slides}:
                page.screenshot(
                    path = str(screenshots_dir / f"slide_{index:02d}.png"),
                    full_page = False
                )
            items = page.evaluate(
                """
                () => {
                  const active = document.querySelector('.slide.is-active');
                  const root = active.getBoundingClientRect();
                  return Array.from(active.querySelectorAll('h1,h2,h3,p,span,footer,.card,.step,.lane,.loop-node,.formula-line,.chart-panel'))
                    .map((el) => {
                      const box = el.getBoundingClientRect();
                      return {
                        tag: el.tagName,
                        text: (el.textContent || '').trim().slice(0, 60),
                        left: box.left,
                        right: box.right,
                        top: box.top,
                        bottom: box.bottom,
                        rootRight: root.right,
                        rootBottom: root.bottom
                      };
                    })
                    .filter((x) => x.right > x.rootRight + 4 || x.bottom > x.rootBottom + 4 || x.left < root.left - 4 || x.top < root.top - 4);
                }
                """
            )
            if items:
                overflow[index] = items[:5]
        browser.close()
    if total != expected_slides:
        raise RuntimeError(f"Expected {expected_slides} rendered HTML slides, found {total}.")
    if overflow:
        raise RuntimeError(f"Detected visible overflow: {overflow}")
    return {
        "slides_rendered": total,
        "screenshots": sorted(path.name for path in screenshots_dir.glob("*.png")),
        "overflow_slides": overflow,
    }


def main():
    """Run deck validation."""
    args = parse_args()
    deck_dir = args.deck_dir.resolve()
    expected_slides = get_expected_slide_count(deck_dir)
    pptx_summary = validate_pptx(deck_dir / PPTX_NAME, expected_slides)
    html_summary = render_html(deck_dir / "index.html", args.screenshots_dir.resolve(), expected_slides)
    logger.info("PPTX summary: %s", pptx_summary)
    logger.info("HTML summary: %s", html_summary)


if __name__ == "__main__":
    logging.basicConfig(
        level = logging.INFO,
        format = '%(asctime)s - %(name)s - %(levelname)s - %(message)s',
        handlers = [logging.StreamHandler()]
    )
    main()

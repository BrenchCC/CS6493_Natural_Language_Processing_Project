import csv
import html
import json
import logging
import argparse
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE
from pptx.chart.data import CategoryChartData
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

logger = logging.getLogger(__name__)


OUT_DIR = Path(__file__).resolve().parent
ROOT_DIR = OUT_DIR.parents[1]
SUMMARY_DIR = ROOT_DIR / "results" / "summaries"
PPTX_NAME = "CS6493_Math_Reasoning_Presentation.pptx"

COLORS = {
    "paper": "F7F9FC",
    "white": "FFFFFF",
    "ink": "18212F",
    "muted": "647084",
    "line": "CED6E3",
    "blue": "1E5B9E",
    "teal": "0F766E",
    "amber": "B7791F",
    "red": "B42318",
    "green": "2F855A",
    "slate": "344256",
    "soft_blue": "EAF2FF",
    "soft_teal": "E7F6F4",
    "soft_amber": "FFF3D6",
    "soft_red": "FEECEB",
    "soft_green": "EAF7EF",
}

METHOD_ORDER = [
    "cot_zero",
    "cot_few_shot",
    "self_refine",
    "self_consistency",
    "plan_solve",
    "tir",
]

METHOD_LABELS = {
    "cot_zero": "CoT zero",
    "cot_few_shot": "CoT few",
    "self_refine": "Self-Refine",
    "self_consistency": "Self-Cons.",
    "plan_solve": "Plan-Solve",
    "tir": "TIR",
}

MODEL_ORDER = [
    "DeepSeek",
    "Qwen",
]

DATASET_ORDER = [
    "math500",
    "gsm8k",
    "aime2024",
]

DATASET_LABELS = {
    "math500": "MATH-500",
    "gsm8k": "GSM8K",
    "aime2024": "AIME 2024",
}


def normalize_model_name(model_alias):
    """Map model aliases from CSV summaries into short slide labels.

    Args:
        model_alias: Raw model alias string from the score summary CSV.

    Returns:
        str: Short display name used in slides.
    """
    if "DeepSeek" in model_alias:
        return "DeepSeek"
    if "Qwen" in model_alias:
        return "Qwen"
    return model_alias


def load_score_rows(path):
    """Load and normalize score summary rows from a CSV file.

    Args:
        path: Path to a summary CSV file.

    Returns:
        list: Normalized row dictionaries with parsed numeric fields.
    """
    float_fields = {
        "accuracy",
        "avg_response_length_tokens",
        "avg_total_response_length_tokens",
        "joint_score",
        "mean_efficiency_on_correct",
        "mean_length_factor_on_correct",
        "mean_answer_factor_on_correct",
        "mean_reflection_factor_on_correct",
    }
    rows = []
    with path.open(encoding = "utf-8") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            normalized = dict(row)
            normalized["sample_count"] = int(row["sample_count"])
            normalized["correct_count"] = int(row["correct_count"])
            normalized["total_count"] = int(row["total_count"])
            normalized["model_display"] = normalize_model_name(row["model_alias"])
            normalized["method_label"] = METHOD_LABELS[row["method_name"]]
            normalized["dataset_label"] = DATASET_LABELS.get(row["dataset_name"], row["dataset_name"])
            for field in float_fields:
                normalized[field] = float(row[field])
            rows.append(normalized)
    return rows


OVERALL_ROWS = load_score_rows(SUMMARY_DIR / "report_selected_overall_scores.csv")
DATASET_ROWS = load_score_rows(SUMMARY_DIR / "report_selected_dataset_scores.csv")

OVERALL_LOOKUP = {
    (row["model_display"], row["method_name"]): row
    for row in OVERALL_ROWS
}

DATASET_LOOKUP = {
    (row["model_display"], row["dataset_name"], row["method_name"]): row
    for row in DATASET_ROWS
}

METHODS = [METHOD_LABELS[method_name] for method_name in METHOD_ORDER]

JOINT_DATA = {
    model_name: [
        OVERALL_LOOKUP[(model_name, method_name)]["joint_score"]
        for method_name in METHOD_ORDER
    ]
    for model_name in MODEL_ORDER
}

FACTOR_DATA = [
    {
        "name": "Answer factor",
        "short": "A_i counts answer cues such as boxed{} or “final answer”. The factor penalizes repeated declarations and long continuation after the first answer cue.",
        "formula_tex": r"a_i = \mathrm{count}_i \cdot \mathrm{tail}_i",
        "formula_ppt": "a_i = count_i · tail_i",
        "color": "blue",
    },
    {
        "name": "Length factor",
        "short": "L_i is the total generated token count over all stages, including plans, critiques, sampled paths, and tool traces. Only over-budget generation is penalized.",
        "formula_tex": r"\ell_i = \exp\left(-\frac{\max(0, L_i - L_{\mathrm{ref}})}{\tau_L}\right)",
        "formula_ppt": "ℓ_i = exp(-max(0, L_i - L_ref) / τ_L)",
        "color": "teal",
    },
    {
        "name": "Reflection factor",
        "short": "F_i counts wait/check/verify/rethink cues. Reflection beyond the first cue is treated as unstable self-correction rather than useful structure.",
        "formula_tex": r"r_i = \exp\left(-\frac{\max(0, F_i - F_{\mathrm{ref}})}{\tau_F}\right)",
        "formula_ppt": "r_i = exp(-max(0, F_i - F_ref) / τ_F)",
        "color": "amber",
    },
]


def build_factor_discussion_cards():
    """Build method-level factor discussion cards from overall summary data.

    Returns:
        list: Card dictionaries for the factor discussion slide.
    """
    sc_deepseek = OVERALL_LOOKUP[("DeepSeek", "self_consistency")]
    sc_qwen = OVERALL_LOOKUP[("Qwen", "self_consistency")]
    plan_deepseek = OVERALL_LOOKUP[("DeepSeek", "plan_solve")]
    plan_qwen = OVERALL_LOOKUP[("Qwen", "plan_solve")]
    tir_deepseek = OVERALL_LOOKUP[("DeepSeek", "tir")]
    tir_qwen = OVERALL_LOOKUP[("Qwen", "tir")]
    return [
        {
            "title": "Self-Consistency buys accuracy with hidden budget",
            "body": (
                f"DeepSeek length factor drops to {sc_deepseek['mean_length_factor_on_correct']:.3f}; "
                f"Qwen still falls to {sc_qwen['mean_length_factor_on_correct']:.3f}. "
                "Five-path voting helps correctness, but the joint score exposes the trajectory cost."
            ),
            "color": "amber",
        },
        {
            "title": "Plan-and-Solve improves by cleaner decomposition",
            "body": (
                f"Reflection stays near one on both models "
                f"({plan_deepseek['mean_reflection_factor_on_correct']:.3f} / "
                f"{plan_qwen['mean_reflection_factor_on_correct']:.3f}). "
                "The gain comes from better early planning rather than extra self-check loops."
            ),
            "color": "blue",
        },
        {
            "title": "TIR is strongest when execution errors matter",
            "body": (
                f"It reaches J = {tir_deepseek['joint_score']:.3f} on DeepSeek and "
                f"J = {tir_qwen['joint_score']:.3f} on Qwen while keeping length factors at "
                f"{tir_deepseek['mean_length_factor_on_correct']:.3f} and "
                f"{tir_qwen['mean_length_factor_on_correct']:.3f}. "
                "Tool use improves execution without collapsing into runaway reasoning."
            ),
            "color": "green",
        },
    ]


def build_model_discussion_cards():
    """Build backbone-level interpretation cards from overall and dataset results.

    Returns:
        list: Card dictionaries for the model discussion slide.
    """
    cot_deepseek = OVERALL_LOOKUP[("DeepSeek", "cot_zero")]
    cot_qwen = OVERALL_LOOKUP[("Qwen", "cot_zero")]
    plan_deepseek = OVERALL_LOOKUP[("DeepSeek", "plan_solve")]
    plan_qwen = OVERALL_LOOKUP[("Qwen", "plan_solve")]
    tir_deepseek = OVERALL_LOOKUP[("DeepSeek", "tir")]
    tir_qwen = OVERALL_LOOKUP[("Qwen", "tir")]
    aime_tir_deepseek = DATASET_LOOKUP[("DeepSeek", "aime2024", "tir")]
    aime_tir_qwen = DATASET_LOOKUP[("Qwen", "aime2024", "tir")]
    return [
        {
            "title": "Both backbones benefit from explicit structure",
            "body": (
                f"Plan-and-Solve adds {plan_deepseek['joint_score'] - cot_deepseek['joint_score']:+.3f} joint "
                f"over CoT zero on DeepSeek and {plan_qwen['joint_score'] - cot_qwen['joint_score']:+.3f} on Qwen. "
                "The decomposition stage helps even without external tools."
            ),
            "color": "blue",
        },
        {
            "title": "A stronger reasoning prior converts tools into bigger hard-split gains",
            "body": (
                f"On AIME 2024, TIR reaches {aime_tir_deepseek['joint_score']:.3f} on DeepSeek "
                f"versus {aime_tir_qwen['joint_score']:.3f} on Qwen. "
                "The hardest split is where backbone quality most clearly limits the return from agentic control."
            ),
            "color": "teal",
        },
        {
            "title": "Interpretation: structure helps both, but not for the same reason",
            "body": (
                f"DeepSeek gets the larger absolute gain with TIR "
                f"({tir_deepseek['joint_score'] - cot_deepseek['joint_score']:+.3f}), "
                f"while Qwen keeps a tighter efficiency profile "
                f"({tir_qwen['mean_length_factor_on_correct']:.3f}). "
                "One model converts extra structure into harder-problem accuracy; the other benefits more from controlled execution."
            ),
            "color": "amber",
        },
    ]


def build_dataset_matrix_panels():
    """Build side-by-side dataset joint-score matrices for both models.

    Returns:
        list: Panel dictionaries for the dataset matrix slide.
    """
    panels = []
    for model_name, accent in [("DeepSeek", "blue"), ("Qwen", "teal")]:
        rows = []
        for method_name in METHOD_ORDER:
            rows.append(
                {
                    "method": METHOD_LABELS[method_name],
                    "scores": [
                        DATASET_LOOKUP[(model_name, dataset_name, method_name)]["joint_score"]
                        for dataset_name in DATASET_ORDER
                    ],
                }
            )
        panels.append(
            {
                "title": f"{model_name} joint score by dataset",
                "accent": accent,
                "rows": rows,
            }
        )
    return panels

SLIDES = [
    {
        "id": 1,
        "kicker": "CS6493 Topic 1",
        "title": "Mathematical Reasoning in Small Open-Weight LLMs",
        "subtitle": "Prompting accuracy and response efficiency for classroom presentation",
        "layout": "cover",
        "tags": [
            "Qwen2.5-Math-1.5B",
            "DeepSeek-R1-Distill-Qwen-1.5B",
            "MATH-500 / GSM8K / AIME 2024",
        ],
        "takeaway": "The project studies not only whether a method gets the answer, but how it behaves while getting there.",
        "notes": [
            "I will present a project on mathematical reasoning in small open-weight language models. The title is intentionally split into two parts: prompting accuracy and response efficiency. The first part is the standard requirement of the course project. The second part is my extension: I want to make the evaluation sensitive to how much generation and unstable answer behavior the method uses.",
            "The deck is organized as a story. I first explain the experimental design, then the prompt methods, then the custom score, and finally the empirical interpretation and future agentic directions. The main claim is that prompt design and metric design should be discussed together, especially once methods become multi-stage or tool-augmented.",
        ],
    },
    {
        "id": 2,
        "kicker": "Research Question",
        "title": "Can structured prompting improve math reasoning under limited compute?",
        "subtitle": "The core question is not just accuracy. It is accuracy under visible reasoning cost.",
        "layout": "cards",
        "cards": [
            {
                "title": "Reasoning quality",
                "body": "Can the model decompose, execute, and return a precise final answer?",
                "color": "blue",
            },
            {
                "title": "Cost visibility",
                "body": "Does the method hide extra generations behind a concise final response?",
                "color": "teal",
            },
            {
                "title": "Behavior stability",
                "body": "Does the answer format stay stable without repeated conclusions or hesitation loops?",
                "color": "amber",
            },
        ],
        "takeaway": "The evaluation target is a reliable solver behavior, not a single isolated number.",
        "notes": [
            "The main research question is whether structured prompting can make small math-oriented models stronger without turning the experiment into an unlimited-compute setting. This matters because the target models are only 1.5B scale, so they do not have the same margin for reasoning errors as larger proprietary models.",
            "I separate the question into three concerns. First, reasoning quality: does the model actually solve the problem? Second, cost visibility: how much text or intermediate generation did it need? Third, behavior stability: does the final answer appear clearly, or does the model keep checking, revising, and repeating itself? This framing leads directly to the custom joint score later in the talk.",
        ],
    },
    {
        "id": 3,
        "kicker": "Experimental Setup",
        "title": "A controlled small-model benchmark across three math splits",
        "subtitle": "Fixed sampled subsets keep the comparison reproducible and affordable.",
        "layout": "setup",
        "cards": [
            {
                "title": "Models",
                "body": "Qwen2.5-Math-1.5B and DeepSeek-R1-Distill-Qwen-1.5B",
                "color": "blue",
            },
            {
                "title": "Datasets",
                "body": "50 MATH-500, 50 GSM8K, and 30 AIME 2024 samples per run",
                "color": "teal",
            },
            {
                "title": "Protocol",
                "body": "Same answer parser, same equivalence grader, best joint-score row selected for reporting",
                "color": "amber",
            },
        ],
        "takeaway": "Each model-method run covers 130 fixed questions, making relative comparisons easier to audit.",
        "notes": [
            "The experiment uses two small open-weight mathematical reasoning models. One is Qwen2.5-Math-1.5B-Instruct, and the other is DeepSeek-R1-Distill-Qwen-1.5B. Both are small enough that prompt structure and generation behavior can make a visible difference.",
            "The data are fixed sampled subsets: 50 questions from MATH-500, 50 from GSM8K, and 30 from AIME 2024. The fixed-sample design makes the project reproducible and affordable, but I also treat it as a caveat because it is not an unbiased full-benchmark estimate. All methods go through the same answer extraction and mathematical equivalence grading pipeline.",
        ],
    },
    {
        "id": 4,
        "kicker": "Experiment Pipeline",
        "title": "The experiment records the full reasoning trajectory, not only the final answer",
        "subtitle": "This is the key design choice that makes multi-stage methods comparable.",
        "layout": "process",
        "steps": [
            {
                "title": "Sample",
                "body": "Use fixed subsets so every method sees the same 130 questions per model.",
                "color": "blue",
            },
            {
                "title": "Infer",
                "body": "Store final response plus intermediate plans, critiques, samples, or tool transcripts.",
                "color": "teal",
            },
            {
                "title": "Grade",
                "body": "Extract final answers with one parser and compare by mathematical equivalence.",
                "color": "amber",
            },
            {
                "title": "Aggregate",
                "body": "Compute joint score and factor diagnostics by model, method, and dataset.",
                "color": "green",
            },
        ],
        "takeaway": "The evaluation unit is the complete solving trajectory, because hidden generations are part of the method.",
        "notes": [
            "The experiment process is deliberately designed around complete trajectories. For a single-turn method, the trajectory is just one response. For a two-turn or multi-turn method, the trajectory includes the plan, critique, refinement, samples, or tool observations. This matters because otherwise multi-stage methods can look artificially cheap.",
            "The pipeline has four steps. First, fixed samples make the comparison reproducible. Second, inference records not only the final answer but also intermediate generated content. Third, the same parser and mathematical equivalence grader are used for all methods. Fourth, aggregation reports the designed joint score and factor diagnostics. This is the bridge between experimental design and metric design.",
        ],
    },
    {
        "id": 5,
        "kicker": "Method Taxonomy",
        "title": "Three intervention levels: prompt context, inference structure, and tool loop",
        "subtitle": "The taxonomy prevents unfair comparison between prompt-only and tool-augmented methods.",
        "layout": "taxonomy",
        "lanes": [
            {
                "title": "Prompt context",
                "methods": [
                    "cot_zero",
                    "cot_few_shot",
                ],
                "body": "Single-turn prompt changes; cleanest assignment-aligned baselines.",
                "color": "blue",
            },
            {
                "title": "Inference structure",
                "methods": [
                    "self_refine",
                    "self_consistency",
                    "plan_solve",
                ],
                "body": "Critique, repeated sampling, or explicit decomposition changes the reasoning trajectory.",
                "color": "teal",
            },
            {
                "title": "Tool loop",
                "methods": [
                    "tir",
                ],
                "body": "The model can emit Python, observe the output, and continue reasoning.",
                "color": "amber",
            },
        ],
        "takeaway": "Similar score gains can come from very different additional computation budgets.",
        "notes": [
            "The six methods are easier to understand if we separate them by intervention level. CoT zero-shot and CoT few-shot mainly change prompt context. Self-Refine, Self-Consistency, and Plan-and-Solve change the structure of the inference process. TIR changes the system boundary because it introduces executable computation.",
            "This taxonomy is important for fair interpretation. A prompt-only method and a tool-augmented loop may both improve the final result, but they are not doing the same thing. The deck will present Plan-and-Solve as the main fair prompt-only extension, while TIR is framed as an agentic contrast case rather than just another baseline.",
        ],
    },
    {
        "id": 6,
        "kicker": "Dialogue Design",
        "title": "Single-turn, multi-sample, multi-turn, and tool-loop methods have different cost surfaces",
        "subtitle": "The dialogue pattern is part of the method, so it must be counted in evaluation.",
        "layout": "method_cards",
        "cards": [
            {
                "title": "Single-turn",
                "body": "CoT zero-shot and few-shot produce one final response; final length equals total length.",
                "color": "blue",
            },
            {
                "title": "Multi-sample",
                "body": "Self-Consistency launches five independent solutions before voting; total cost grows linearly.",
                "color": "teal",
            },
            {
                "title": "Multi-turn",
                "body": "Self-Refine and Plan-and-Solve split reasoning into stages, so intermediate text must be counted.",
                "color": "amber",
            },
            {
                "title": "Tool-loop",
                "body": "TIR inserts Python actions and observations; the transcript becomes part of the solving cost.",
                "color": "green",
            },
        ],
        "takeaway": "A fair comparison needs both final answer quality and full trajectory accounting.",
        "notes": [
            "This slide makes the single-round and multi-round design explicit. CoT zero-shot and few-shot are single-turn methods, so their final response length and total generated length are the same. Self-Consistency is not a dialogue, but it is multi-sample because it generates five independent solutions before voting.",
            "Self-Refine and Plan-and-Solve are multi-turn or multi-stage designs. Their final answers can be short, but the full trajectory contains earlier solution, critique, plan, or solve stages. TIR goes one step further by adding executable actions and tool observations. Because the dialogue pattern changes the cost surface, it must be part of the evaluation design.",
        ],
    },
    {
        "id": 7,
        "kicker": "Prompt Design Problem",
        "title": "Small reasoning models fail in more than one way",
        "subtitle": "A wrong answer is only the visible end of several behavioral failure modes.",
        "layout": "cards",
        "cards": [
            {
                "title": "Missing-step errors",
                "body": "The model starts solving before identifying all required operations.",
                "color": "blue",
            },
            {
                "title": "Overlong reasoning",
                "body": "The model keeps generating even after the useful reasoning signal has peaked.",
                "color": "teal",
            },
            {
                "title": "Unstable final answers",
                "body": "The model repeats answer cues, revises late, or continues after the answer.",
                "color": "red",
            },
        ],
        "takeaway": "The prompt design should reduce planning errors, and the metric should expose inefficient reasoning behavior.",
        "notes": [
            "Before introducing the methods, I want to state the prompt design problem. In mathematical reasoning, a failure is not only a wrong final answer. The model can miss an important step at the beginning, produce an extremely long derivation, or keep changing the final answer signal near the end.",
            "These behaviors matter because different prompts address different causes. Decomposition helps when the model starts too quickly. Sampling helps when one reasoning path is unreliable. Tool use helps when arithmetic or symbolic execution is the bottleneck. This is why the project needs both method diversity and a metric that can diagnose behavior beyond correctness.",
        ],
    },
    {
        "id": 8,
        "kicker": "Method Characteristics",
        "title": "Each method changes a different part of the reasoning process",
        "subtitle": "The methods are not only prompt names; each one implies a different dialogue and compute pattern.",
        "layout": "method_cards",
        "cards": [
            {
                "title": "CoT zero-shot",
                "body": "One response. Tests whether step-by-step instruction alone can trigger useful reasoning.",
                "color": "blue",
            },
            {
                "title": "CoT few-shot",
                "body": "One response with demonstrations. Guides style, but examples can also bias format.",
                "color": "teal",
            },
            {
                "title": "Self-Refine",
                "body": "Solve, critique, rewrite. Useful when the model can detect its own mistakes.",
                "color": "amber",
            },
            {
                "title": "Self-Consistency",
                "body": "Five independent paths plus voting. Robust but the most expensive baseline.",
                "color": "green",
            },
            {
                "title": "Plan-and-Solve",
                "body": "Plan first, solve second. Targets missing-step errors without external tools.",
                "color": "blue",
            },
            {
                "title": "TIR",
                "body": "Reason, call Python, observe, continue. Strong but tool-augmented rather than prompt-only.",
                "color": "teal",
            },
        ],
        "takeaway": "The most important contrast is not baseline versus extension, but where extra reasoning budget enters the system.",
        "notes": [
            "Here I summarize the six methods as behavioral designs rather than only names. CoT zero-shot tests a pure instruction. CoT few-shot adds demonstrations. Self-Refine asks the model to critique and revise itself. Self-Consistency uses repeated sampling and voting to improve robustness.",
            "Plan-and-Solve is different because it inserts structure before the solution. TIR is different again because it allows a tool observation to enter the reasoning loop. This slide is useful for the audience because it explains why the methods cannot be evaluated only by a final answer string. Each method changes the amount and location of extra reasoning budget.",
        ],
    },
    {
        "id": 7,
        "kicker": "Prompt Innovation",
        "title": "Plan-and-Solve moves structure before execution",
        "subtitle": "The proposed prompt-only method asks the model to plan first, then solve according to the plan.",
        "layout": "process",
        "steps": [
            {
                "title": "Plan",
                "body": "Write a short numbered decomposition of the mathematical task.",
                "color": "blue",
            },
            {
                "title": "Solve",
                "body": "Follow the plan step by step, avoiding unsupported jumps.",
                "color": "teal",
            },
            {
                "title": "Finalize",
                "body": "Return one clear final answer for parsing and grading.",
                "color": "amber",
            },
        ],
        "takeaway": "Plan-and-Solve is fair against prompt-only baselines because it adds structure without external tools.",
        "notes": [
            "Plan-and-Solve is the main prompt-only extension. The key design choice is timing. Instead of asking the model to refine after a full solution, it asks for a plan before the solution. This is a better fit for math problems where one missing early step can derail the entire derivation.",
            "It is also easier to justify under limited compute. Unlike Self-Consistency, it does not launch five independent solutions. Unlike TIR, it does not call an external Python executor. The extra cost is conceptually transparent: one planning stage and one solving stage. That makes it a strong classroom story as well as a practical prompt design.",
        ],
    },
    {
        "id": 8,
        "kicker": "Agentic Extension",
        "title": "TIR turns math solving into a compact action-observation loop",
        "subtitle": "Tool-Integrated Reasoning is not prompt-only; it is a small agentic workflow.",
        "layout": "loop",
        "steps": [
            {
                "title": "Reason",
                "body": "Identify whether the next step needs reliable computation.",
                "color": "blue",
            },
            {
                "title": "Act",
                "body": "Emit executable Python when arithmetic or symbolic work is useful.",
                "color": "teal",
            },
            {
                "title": "Observe",
                "body": "Read the tool output as evidence for the next reasoning step.",
                "color": "amber",
            },
            {
                "title": "Continue",
                "body": "Update the solution and produce a final answer.",
                "color": "green",
            },
        ],
        "takeaway": "TIR is closest to a ReAct-style solver loop, but with a narrow mathematical tool boundary.",
        "notes": [
            "TIR is the most distinctive method, but it needs careful framing. It should not be presented as another pure prompt baseline because it changes the setting. The model can emit Python code, observe the returned output, and continue solving with that evidence.",
            "This is why I describe it as a compact agentic pattern. It resembles the reason-act-observe style of agent loops, but the scope is deliberately narrow: mathematical computation, not open-ended web browsing or autonomous task execution. The research value is that it separates reasoning-plan errors from execution errors, especially for arithmetic-heavy questions.",
        ],
    },
    {
        "id": 9,
        "kicker": "Evaluation Motivation",
        "title": "Single metrics hide different failure modes",
        "subtitle": "Once methods become multi-stage, accuracy and length stop telling the same story.",
        "layout": "cards",
        "cards": [
            {
                "title": "Accuracy-first gate",
                "body": "Wrong answers receive zero, so the score never rewards a short but incorrect response.",
                "color": "blue",
            },
            {
                "title": "Full trajectory cost",
                "body": "Plans, critiques, five samples, and tool transcripts are counted as generated reasoning cost.",
                "color": "teal",
            },
            {
                "title": "Behavior diagnostics",
                "body": "Answer, length, and reflection factors explain why two correct methods differ.",
                "color": "amber",
            },
        ],
        "takeaway": "The metric is designed to be accuracy-first, trajectory-aware, and interpretable during error analysis.",
        "notes": [
            "This slide explains why I do not want to present raw accuracy, final answer length, or total generation length as isolated scoreboards. Accuracy alone hides hidden budget. Final length hides the cost of multi-stage trajectories. Total length alone can over-penalize correct but slightly verbose reasoning.",
            "The joint score is therefore a compromise with a clear priority order. Correctness decides whether the sample deserves credit at all. Full trajectory accounting makes multi-stage methods comparable. Then answer, length, and reflection factors explain why two correct methods still differ in quality.",
        ],
    },
    {
        "id": 10,
        "kicker": "Designed Score Formula",
        "title": "Correctness gates the score; efficiency adjusts the remaining credit",
        "subtitle": "The formula itself is part of the contribution, so it is rendered explicitly instead of being described only in prose.",
        "layout": "formula",
        "formula": [
            {
                "tex": r"\eta_i = a_i^{0.5}\,\ell_i^{0.3}\,r_i^{0.2}",
                "ppt": "η_i = a_i^0.5 · ℓ_i^0.3 · r_i^0.2",
                "desc": "Efficiency is a weighted geometric mean of answer stability, length economy, and reflection control.",
            },
            {
                "tex": r"s_i = c_i\left(0.6 + 0.4\eta_i\right), \quad c_i \in \{0,1\}",
                "ppt": "s_i = c_i(0.6 + 0.4η_i),  c_i ∈ {0, 1}",
                "desc": "Correctness gates the score: incorrect predictions receive zero regardless of brevity.",
            },
            {
                "tex": r"J = \frac{1}{N}\sum_{i = 1}^{N} s_i",
                "ppt": "J = (1 / N) · Σ_i=1^N s_i",
                "desc": "The reported joint score averages sample-level scores over a model-method run.",
            },
        ],
        "definitions": [
            {"tex": r"A_i", "ppt": "A_i", "desc": "answer-cue count"},
            {"tex": r"R_i", "ppt": "R_i", "desc": "tail ratio after first cue"},
            {"tex": r"L_i", "ppt": "L_i", "desc": "total generated tokens"},
            {"tex": r"F_i", "ppt": "F_i", "desc": "reflection-cue count"},
            {"tex": r"L_{\mathrm{ref}} = 1024,\ \tau_L = 512", "ppt": "L_ref = 1024, τ_L = 512", "desc": "length reference"},
            {"tex": r"F_{\mathrm{ref}} = 1,\ \tau_F = 2", "ppt": "F_ref = 1, τ_F = 2", "desc": "reflection reference"},
        ],
        "takeaway": "The 0.6 floor keeps correctness dominant, while the 0.4 modifier separates correct samples by behavior.",
        "notes": [
            "This is the core scoring slide. I make the formula explicit because the evaluation design is not just a reporting afterthought; it is part of the experiment. The exponents tell the audience that answer stability is the strongest behavior term, followed by length and then reflection.",
            "The second equation is the key design decision. Correctness gates the entire score, and efficiency only adjusts the remaining forty percent. That makes the metric conservative: a short wrong answer still receives zero, but among correct samples we can distinguish cleaner reasoning from wasteful reasoning.",
        ],
    },
    {
        "id": 11,
        "kicker": "Factor Decomposition",
        "title": "Factor and variable design make the score interpretable",
        "subtitle": "Each penalty is tied to a concrete observable variable instead of an opaque heuristic.",
        "layout": "factors",
        "cards": FACTOR_DATA,
        "takeaway": "Low length, answer, or reflection factors point to different error-analysis stories.",
        "notes": [
            "The important point here is that each factor is linked to a variable I can actually observe in the trajectory. A_i and R_i describe answer behavior. L_i describes total budget across all generated stages. F_i describes unstable reflection language such as wait, check, or verify.",
            "That design makes the score explainable. If a method loses score, I can say whether it is because it over-generated, repeated final-answer cues, or kept looping through self-check language. This is much more useful in discussion than saying the method simply had a lower scalar score.",
        ],
    },
    {
        "id": 12,
        "kicker": "Designed Score View",
        "title": "Joint score comparison highlights the proposed methods",
        "subtitle": "The result slide uses only the custom accuracy-first joint score, not raw accuracy or raw length charts.",
        "layout": "joint_chart",
        "takeaway": "Plan-and-Solve leads prompt-only methods; TIR is strongest when tool use is allowed.",
        "notes": [
            "This is the main result view, and it intentionally uses only joint score. The audience should read the chart as a ranking under my full evaluation design, not as a replacement for the underlying correctness task.",
            "The result pattern is stable across both models. Plan-and-Solve is the strongest prompt-only method. TIR is best overall when tool use is allowed. The substantive interpretation is that structure helps small models, and action-observation loops help even more when execution errors are the bottleneck.",
        ],
    },
    {
        "id": 13,
        "kicker": "Factor-Based Discussion",
        "title": "Factor diagnostics separate efficient gains from expensive gains",
        "subtitle": "Methods can improve correctness for very different behavioral reasons; the factor view reveals which one.",
        "layout": "cards",
        "cards": build_factor_discussion_cards(),
        "takeaway": "The factor view explains whether a gain comes from cleaner structure, more budget, or better execution.",
        "notes": [
            "This slide is where the factor table becomes an argument rather than a definition. Self-Consistency is a good example: it may improve correctness, but its length factor collapses because five full trajectories are being paid for. The joint score makes that trade-off visible.",
            "Plan-and-Solve and TIR tell a different story. Plan-and-Solve looks strong because planning reduces unstable reflection while staying prompt-only. TIR looks strong because it improves execution reliability while keeping the trajectory more controlled than multi-sample voting.",
        ],
    },
    {
        "id": 14,
        "kicker": "Model-Level Discussion",
        "title": "Backbone strength changes how much structure can be exploited",
        "subtitle": "Method gains are real on both models, but the stronger reasoning prior converts them more effectively on hard cases.",
        "layout": "cards",
        "cards": build_model_discussion_cards(),
        "takeaway": "Structure helps both backbones; stronger priors turn it into bigger hard-split gains.",
        "notes": [
            "Here I shift from method comparison to backbone interpretation. The key point is that both models benefit from structure, which means the prompt idea is not tied to one particular backbone. But the gain profile is not identical.",
            "DeepSeek converts extra structure and tool access into larger gains on the hardest split, especially AIME. Qwen also improves, but its ceiling is lower. So the takeaway is not only that prompting matters, but that a stronger reasoning prior can better exploit the same prompting or agentic scaffold.",
        ],
    },
    {
        "id": 15,
        "kicker": "Dataset-Level Discussion",
        "title": "Dataset joint-score matrix shows where the gains really transfer",
        "subtitle": "GSM8K reveals decomposition gains early; AIME 2024 checks whether those gains survive on harder symbolic reasoning.",
        "layout": "dataset_matrix",
        "panels": build_dataset_matrix_panels(),
        "takeaway": "GSM8K responds fastest to structure; AIME remains the transfer stress test for both backbones.",
        "notes": [
            "This matrix is still using only joint score, so it stays inside the evaluation design of the project. The purpose is not to create another scoreboard, but to show where a method’s gains actually transfer across data regimes.",
            "GSM8K is where decomposition and tools help most predictably, because many questions are arithmetic and multi-step. MATH-500 is broader and more mixed. AIME 2024 is the hardest test: if a method only improves answer style or easy arithmetic, its joint score will collapse there first.",
        ],
    },
    {
        "id": 16,
        "kicker": "Literature Connections",
        "title": "The design borrows ideas from prompting, tool use, and reasoning efficiency",
        "subtitle": "The references are used as design motivation, not as claims of full system equivalence.",
        "layout": "literature",
        "cards": [
            {
                "title": "Prompting lineage",
                "body": "Wei et al. (2022), Wang et al. (2022), Madaan et al. (2024), and Wang et al. (2023) motivate elicitation, sampling, critique, and explicit planning as distinct reasoning scaffolds.",
                "color": "blue",
            },
            {
                "title": "Tool-grounded solving",
                "body": "ToRA motivates the idea that reasoning can call reliable computation, observe the result, and continue. My TIR slide is a prompt-level, small-scale version of that idea.",
                "color": "teal",
            },
            {
                "title": "Efficiency-aware diagnosis",
                "body": "NoWait and Dynamic Early Exit inspire the metric side: reflection tokens, overlong reasoning, and answer-aware stopping should be measured instead of ignored.",
                "color": "amber",
            },
        ],
        "takeaway": "The project adapts these ideas into a lightweight prompt-level experimental pipeline.",
        "notes": [
            "This slide clarifies that the project is not claiming to reproduce full prior systems. Instead, it borrows the key ideas that matter for this classroom experiment: elicited reasoning, multi-path robustness, pre-solution planning, tool-grounded execution, and efficiency-aware analysis.",
            "That framing matters because it ties each design choice back to a paper idea. Plan-and-Solve is not just a prompt trick; it is motivated by explicit task decomposition. TIR is not just longer prompting; it is motivated by tool-grounded reasoning. The factorized score is not arbitrary; it is motivated by recent work on thinking-token cost and stopping decisions.",
        ],
    },
    {
        "id": 17,
        "kicker": "Future Work",
        "title": "From static prompts to adaptive solver agents",
        "subtitle": "The next step is not “always think longer”, but “decide which reasoning mode is worth paying for”.",
        "layout": "roadmap",
        "steps": [
            {
                "title": "Harness prompt",
                "body": "Standardize classify-plan-solve-verify slots so later agents read and act on the same reasoning scaffold.",
                "color": "blue",
            },
            {
                "title": "Problem router",
                "body": "Choose CoT, Plan-and-Solve, or TIR according to arithmetic load, symbolic depth, and uncertainty.",
                "color": "teal",
            },
            {
                "title": "ReAct math loop",
                "body": "Let the solver reason, call Python, observe, and revise only when execution evidence is actually needed.",
                "color": "amber",
            },
            {
                "title": "Verifier + stop rule",
                "body": "Check final-answer stability, repeated cues, and contradiction; then stop early or request one targeted repair.",
                "color": "green",
            },
        ],
        "takeaway": "A practical next experiment is a routed solver: classify the problem, choose the reasoning mode, verify, then stop early when stable.",
        "notes": [
            "The future-work direction is intentionally agentic, but still grounded in this project. A harness prompt would give every method a shared scaffold with slots for classification, planning, solving, verification, and stopping. That makes later control easier.",
            "Then a router decides whether the cheapest sufficient mode is plain CoT, Plan-and-Solve, or a TIR-style loop. A verifier agent checks final-answer stability and contradiction. The system stops early when the answer is stable and only pays for more reasoning when the evidence says it should.",
        ],
    },
    {
        "id": 18,
        "kicker": "Takeaways",
        "title": "Prompt design and metric design must be discussed together",
        "subtitle": "Once reasoning becomes multi-stage or tool-augmented, evaluation must describe the behavior it rewards.",
        "layout": "takeaways",
        "cards": [
            {
                "title": "1",
                "body": "Plan-and-Solve is the cleanest main contribution because it improves joint score without leaving the prompt-only setting.",
                "color": "blue",
            },
            {
                "title": "2",
                "body": "TIR is best read as an agentic extension: reason, act, observe, and continue with a narrow mathematical tool boundary.",
                "color": "teal",
            },
            {
                "title": "3",
                "body": "The joint score keeps correctness dominant while exposing hidden budget, unstable answer cues, and excessive reflection.",
                "color": "amber",
            },
        ],
        "takeaway": "Q&A",
        "notes": [
            "The closing message has three parts. First, Plan-and-Solve is the strongest prompt-only story because it is fair, structured, and empirically strong under the designed score. Second, TIR points toward agentic math solvers rather than just longer prompts.",
            "Third, metric design and prompt design should be discussed together. Once a method uses planning, sampling, refinement, or tools, evaluation has to say what budget it counts and what behavior it rewards. That is the methodological takeaway I want the audience to remember.",
        ],
    },
]


EXPLANATIONS = {
    "Mathematical Reasoning in Small Open-Weight LLMs": "This presentation studies mathematical reasoning as a complete solving behavior. The main contribution is a fair comparison of prompt-only designs and a behavior-aware score that makes hidden reasoning cost explicit.",
    "Can structured prompting improve math reasoning under limited compute?": "The research question is framed around small open-weight models, where each extra generation stage matters. The goal is to improve correctness while keeping the reasoning process observable and comparable.",
    "A controlled small-model benchmark across three math splits": "The setup uses fixed sampled subsets to make the experiment reproducible. This supports relative method comparison, while the report still treats exact numbers as sample-dependent reference values.",
    "The experiment records the full reasoning trajectory, not only the final answer": "For multi-stage methods, the intermediate plan, critique, voted samples, and tool transcript are part of the method. Counting them avoids underestimating the real computational behavior.",
    "Three intervention levels: prompt context, inference structure, and tool loop": "The taxonomy separates methods by what they change: the prompt text, the inference trajectory, or the system boundary. This prevents treating tool-augmented gains as ordinary prompt-only gains.",
    "Single-turn, multi-sample, multi-turn, and tool-loop methods have different cost surfaces": "A single response, five sampled responses, staged dialogue, and action-observation loop all expose different costs. The evaluation therefore uses total generated length rather than only final answer length.",
    "Small reasoning models fail in more than one way": "The prompt design targets failure modes beyond final wrongness: missing early steps, excessive derivation, and unstable answer formatting. These behaviors motivate both Plan-and-Solve and the factorized metric.",
    "Each method changes a different part of the reasoning process": "The six methods represent distinct mechanisms: instruction, demonstration, self-critique, sampling, pre-solution planning, and external computation. This makes method interpretation more important than just ranking.",
    "Plan-and-Solve moves structure before execution": "Plan-and-Solve is the main prompt-only innovation because it encourages problem decomposition before arithmetic execution. It is easy to compare fairly against CoT and refinement baselines.",
    "TIR turns math solving into a compact action-observation loop": "TIR is framed as a small agentic extension. Its value is not merely higher score, but the ability to externalize exact computation and feed the observation back into reasoning.",
    "Single metrics hide different failure modes": "Accuracy, final length, and total length each miss part of the story when used alone. The designed metric keeps accuracy central while exposing the cost and stability of the reasoning path.",
    "Correctness gates the score; efficiency adjusts the remaining credit": "The formula is deliberately conservative: correctness is a prerequisite, then behavior factors adjust only the remaining credit. This protects against rewarding short but wrong outputs.",
    "Factor and variable design make the score interpretable": "The factor decomposition makes the score interpretable. A low score can be traced to over-generation, repeated answer cues, or excessive reflection rather than treated as an opaque penalty.",
    "Joint score comparison highlights the proposed methods": "The result view focuses on the designed joint score because it is the evaluation contribution of the project. Plan-and-Solve leads prompt-only methods, while TIR shows the potential of tool-augmented solving.",
    "Factor diagnostics separate efficient gains from expensive gains": "The factor view reveals whether a score gain comes from cleaner decomposition, better execution, or simply paying for more hidden generation budget.",
    "Backbone strength changes how much structure can be exploited": "Method gains depend on model quality as well as prompt design. The stronger backbone converts the same structured scaffold into larger hard-split gains.",
    "Dataset joint-score matrix shows where the gains really transfer": "GSM8K reveals early decomposition gains, MATH-500 broadens coverage, and AIME tests whether the gain survives under harder symbolic reasoning.",
    "The design borrows ideas from prompting, tool use, and reasoning efficiency": "The related work motivates the design choices: CoT-style reasoning paths, ToRA-style tool integration, and recent efficiency work on reflection and overlong reasoning.",
    "From static prompts to adaptive solver agents": "A concrete next step is an adaptive solver that routes problems to the cheapest sufficient reasoning mode, verifies answer stability, and stops once the solution is reliable.",
    "Prompt design and metric design must be discussed together": "The final message is methodological: once methods use planning, sampling, refinement, or tools, evaluation must specify what behavior it rewards and what cost it counts.",
}


for slide_index, slide_data in enumerate(SLIDES, start = 1):
    slide_data["id"] = slide_index
    slide_data["explanation"] = EXPLANATIONS.get(slide_data["title"], slide_data["takeaway"])


def parse_args():
    """Parse command-line arguments.

    Returns:
        argparse.Namespace: Parsed command-line options.
    """
    parser = argparse.ArgumentParser(
        description = "Build the CS6493 classroom HTML deck and editable PPTX."
    )
    parser.add_argument(
        "--out-dir",
        type = Path,
        default = OUT_DIR,
        help = "Directory where deck files should be written."
    )
    parser.add_argument(
        "--skip-pptx",
        action = "store_true",
        help = "Write HTML assets only and skip PPTX generation."
    )
    return parser.parse_args()


def rgb(hex_color):
    """Convert a hex color string into an RGBColor object.

    Args:
        hex_color: Six-character RGB hex string, with or without a leading hash.

    Returns:
        RGBColor: PowerPoint-compatible RGB color value.
    """
    value = hex_color.strip().lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def color(name):
    """Return a hex color by palette name.

    Args:
        name: Key in the shared COLORS palette.

    Returns:
        str: Six-character RGB hex string.
    """
    return COLORS[name]


def escape_text(text):
    """Escape text for safe HTML output.

    Args:
        text: Raw text content.

    Returns:
        str: HTML-escaped text.
    """
    return html.escape(str(text), quote = True)


def standard_formula(text):
    """Convert internal ASCII formula notation to presentation notation.

    Args:
        text: Formula text using ASCII-friendly names.

    Returns:
        str: Formula text with mathematical symbols for slides.
    """
    replacements = {
        "eta_i": "ηᵢ",
        "a_i": "aᵢ",
        "ell_i": "ℓᵢ",
        "r_i": "rᵢ",
        "s_i": "sᵢ",
        "c_i": "cᵢ",
        "L_i": "Lᵢ",
        "F_i": "Fᵢ",
        "L_ref": "L_ref",
        "F_ref": "F_ref",
        "tau_L": "τ_L",
        "tau_F": "τ_F",
        "sum_{i=1}^{N}": "∑ᵢ₌₁ᴺ",
        " in ": " ∈ ",
        " x ": " · ",
    }
    result = str(text)
    for old, new in replacements.items():
        result = result.replace(old, new)
    return result


def write_text(path, content):
    """Write UTF-8 text to a file.

    Args:
        path: Destination file path.
        content: Text content to write.
    """
    path.parent.mkdir(parents = True, exist_ok = True)
    path.write_text(content, encoding = "utf-8")


def render_chip(text):
    """Render a small HTML chip.

    Args:
        text: Chip label.

    Returns:
        str: HTML string for the chip.
    """
    return f'<span class="chip">{escape_text(text)}</span>'


def render_card(card):
    """Render a content card for the HTML deck.

    Args:
        card: Card dictionary with title, body, and color fields.

    Returns:
        str: HTML string for one card.
    """
    card_color = card.get("color", "blue")
    title = card.get("title", card.get("name", ""))
    body = card.get("body", card.get("short", ""))
    formula = card.get("formula")
    formula_tex = card.get("formula_tex")
    body_html = escape_text(body)
    if formula_tex:
        body_html = (
            f'{body_html}<span class="mini-equation math-inline" data-math = "1">'
            f'\\({escape_text(formula_tex)}\\)'
            f'</span>'
        )
    elif formula:
        body_html = f'{body_html}<span class="mini-equation">{escape_text(standard_formula(formula))}</span>'
    return (
        f'<article class="card accent-{card_color}">'
        f'<h3>{escape_text(title)}</h3>'
        f'<p>{body_html}</p>'
        f'</article>'
    )


def render_notes(slide):
    """Render hidden speaker notes for one HTML slide.

    Args:
        slide: Slide dictionary containing a notes list.

    Returns:
        str: HTML string for the notes block.
    """
    paragraphs = "".join(
        f"<p>{escape_text(paragraph)}</p>"
        for paragraph in slide.get("notes", [])
    )
    return f'<aside class="notes">{paragraphs}</aside>'


def get_explanation_meta(slide):
    """Return a label, meta tag, and accent for the explanation module.

    Args:
        slide: Slide dictionary containing id and layout fields.

    Returns:
        tuple: (label, meta tag, accent color key)
    """
    by_id = {
        2: ("Research Lens", "QUESTION", "blue"),
        7: ("Failure Lens", "PROMPT", "amber"),
        11: ("Metric Logic", "EVALUATION", "teal"),
        12: ("Score Rule", "FORMULA", "blue"),
        13: ("Factor Read", "DIAGNOSTIC", "teal"),
        14: ("Result Read", "OVERALL", "blue"),
        15: ("Factor Read", "DIAGNOSTIC", "green"),
        16: ("Model Lens", "BACKBONE", "blue"),
        17: ("Dataset Lens", "JOINT SCORE", "teal"),
        18: ("Reference Link", "RELATED WORK", "blue"),
        19: ("Next Move", "FUTURE", "amber"),
        20: ("Final Takeaway", "Q&A", "green"),
    }
    by_layout = {
        "cover": ("Core Claim", "PROJECT", "blue"),
        "setup": ("Protocol Read", "SETUP", "teal"),
        "taxonomy": ("Comparison Rule", "METHOD", "amber"),
        "process": ("Design Lens", "PIPELINE", "blue"),
        "method_cards": ("Method Read", "BEHAVIOR", "teal"),
        "cards": ("Key Read", "FRAMING", "blue"),
        "loop": ("Agentic Angle", "REASONING", "amber"),
        "formula": ("Score Rule", "FORMULA", "blue"),
        "factors": ("Factor Read", "DIAGNOSTIC", "teal"),
        "joint_chart": ("Result Read", "RESULT", "blue"),
        "comparison": ("Contrast Read", "RESULT", "amber"),
        "dataset_cards": ("Dataset Lens", "SPLITS", "teal"),
        "dataset_matrix": ("Dataset Lens", "JOINT SCORE", "teal"),
        "literature": ("Reference Link", "RELATED WORK", "blue"),
        "roadmap": ("Next Move", "FUTURE", "amber"),
        "takeaways": ("Closing Note", "Q&A", "green"),
    }
    if slide["id"] in by_id:
        return by_id[slide["id"]]
    return by_layout.get(slide["layout"], ("Key Read", "DECK", "blue"))


def render_explanation(slide):
    """Render the visible academic explanation box for one slide.

    Args:
        slide: Slide dictionary containing an explanation string.

    Returns:
        str: HTML markup for the explanation box.
    """
    explanation = slide.get("explanation", "")
    if not explanation:
        return ""
    label, meta, accent = get_explanation_meta(slide)
    return (
        f'<div class="explain-box accent-{accent}">'
        f'<div class="explain-head">'
        f'<span class="explain-tag">{escape_text(label)}</span>'
        f'<span class="explain-meta">{escape_text(meta)}</span>'
        f'</div>'
        f'<p>{escape_text(explanation)}</p>'
        f'</div>'
    )


def score_band(score):
    """Map a joint-score value to a CSS band label.

    Args:
        score: Joint-score value between zero and one.

    Returns:
        str: Band label used by HTML table cells.
    """
    if score >= 0.75:
        return "top"
    if score >= 0.60:
        return "high"
    if score >= 0.35:
        return "mid"
    if score >= 0.20:
        return "low"
    return "risk"


def render_joint_svg():
    """Render an editable-looking SVG bar chart for the HTML deck.

    Returns:
        str: Inline SVG chart markup for joint score comparison.
    """
    max_value = 0.75
    row_gap = 44
    start_y = 70
    label_x = 18
    bar_x = 178
    chart_width = 560
    rows = []
    for index, method in enumerate(METHODS):
        y = start_y + index * row_gap
        ds_width = int(JOINT_DATA["DeepSeek"][index] / max_value * chart_width)
        qw_width = int(JOINT_DATA["Qwen"][index] / max_value * chart_width)
        rows.append(
            f'<text x="{label_x}" y="{y + 17}" class="axis">{escape_text(method)}</text>'
            f'<rect x="{bar_x}" y="{y}" width="{ds_width}" height="14" rx="3" class="bar ds"/>'
            f'<rect x="{bar_x}" y="{y + 18}" width="{qw_width}" height="14" rx="3" class="bar qw"/>'
            f'<text x="{bar_x + ds_width + 8}" y="{y + 12}" class="value">{JOINT_DATA["DeepSeek"][index]:.3f}</text>'
            f'<text x="{bar_x + qw_width + 8}" y="{y + 30}" class="value">{JOINT_DATA["Qwen"][index]:.3f}</text>'
        )
    return (
        '<div class="chart-panel">'
        '<div class="legend"><span class="dot ds"></span>DeepSeek joint score '
        '<span class="dot qw"></span>Qwen joint score</div>'
        '<svg class="joint-svg" viewBox="0 0 820 360" role="img" aria-label="Joint score chart">'
        '<line x1="178" y1="42" x2="738" y2="42" class="gridline"/>'
        '<text x="178" y="28" class="axis">0.000</text>'
        '<text x="720" y="28" class="axis">0.750</text>'
        f'{"".join(rows)}'
        '</svg>'
        '</div>'
    )


def render_dataset_matrix(slide):
    """Render the dataset joint-score matrix for the HTML deck.

    Args:
        slide: Slide dictionary containing dataset matrix panels.

    Returns:
        str: HTML markup for the dataset matrix slide.
    """
    panels = []
    headers = "".join(f"<th>{escape_text(DATASET_LABELS[name])}</th>" for name in DATASET_ORDER)
    for panel in slide.get("panels", []):
        body_rows = []
        for row in panel["rows"]:
            score_cells = "".join(
                (
                    f'<td class="score-cell band-{score_band(score)}">'
                    f'{score:.3f}'
                    f'</td>'
                )
                for score in row["scores"]
            )
            body_rows.append(
                f'<tr>'
                f'<th>{escape_text(row["method"])}</th>'
                f'{score_cells}'
                f'</tr>'
            )
        panels.append(
            f'<article class="matrix-panel accent-{panel["accent"]}">'
            f'<h3>{escape_text(panel["title"])}</h3>'
            f'<table class="score-table">'
            f'<thead><tr><th>Method</th>{headers}</tr></thead>'
            f'<tbody>{"".join(body_rows)}</tbody>'
            f'</table>'
            f'</article>'
        )
    return f'<div class="dataset-matrix">{"".join(panels)}</div>'


def render_visual(slide):
    """Render the main visual area for one HTML slide.

    Args:
        slide: Slide dictionary with layout-specific content.

    Returns:
        str: HTML markup for the slide body.
    """
    layout = slide["layout"]
    if layout == "cover":
        chips = "".join(render_chip(tag) for tag in slide.get("tags", []))
        return (
            f'<div class="cover-layout">'
            f'<div>'
            f'<p class="kicker">{escape_text(slide["kicker"])}</p>'
            f'<h1>{escape_text(slide["title"])}</h1>'
            f'<p class="subtitle">{escape_text(slide["subtitle"])}</p>'
            f'<div class="chips">{chips}</div>'
            f'</div>'
            f'<div class="cover-mark"><span>{len(SLIDES)}</span><small>slides</small></div>'
            f'</div>'
        )
    if layout in {"cards", "method_cards", "dataset_cards"}:
        cards = "".join(render_card(card) for card in slide.get("cards", []))
        return f'<div class="card-grid count-{len(slide.get("cards", []))}">{cards}</div>'
    if layout == "setup":
        cards = "".join(render_card(card) for card in slide.get("cards", []))
        return (
            f'<div class="setup-grid">{cards}</div>'
            f'<div class="callout">Total per model-method run: 130 fixed questions.</div>'
        )
    if layout == "taxonomy":
        lanes = []
        for lane in slide.get("lanes", []):
            methods = "".join(render_chip(method) for method in lane["methods"])
            lanes.append(
                f'<article class="lane accent-{lane["color"]}">'
                f'<h3>{escape_text(lane["title"])}</h3>'
                f'<div class="chips">{methods}</div>'
                f'<p>{escape_text(lane["body"])}</p>'
                f'</article>'
            )
        return f'<div class="lane-grid">{"".join(lanes)}</div>'
    if layout in {"process", "roadmap"}:
        steps = []
        for index, step in enumerate(slide.get("steps", []), start = 1):
            steps.append(
                f'<article class="step accent-{step["color"]}">'
                f'<span>{index:02d}</span>'
                f'<h3>{escape_text(step["title"])}</h3>'
                f'<p>{escape_text(step["body"])}</p>'
                f'</article>'
            )
        return f'<div class="step-grid count-{len(slide.get("steps", []))}">{"".join(steps)}</div>'
    if layout == "loop":
        steps = []
        for index, step in enumerate(slide.get("steps", []), start = 1):
            arrow = " -> " if index < len(slide.get("steps", [])) else ""
            steps.append(
                f'<article class="loop-node accent-{step["color"]}">'
                f'<span>{index}</span>'
                f'<h3>{escape_text(step["title"])}</h3>'
                f'<p>{escape_text(step["body"])}</p>'
                f'</article>{f"<b>{arrow}</b>" if arrow else ""}'
            )
        return f'<div class="loop-grid">{"".join(steps)}</div>'
    if layout == "formula":
        formulas = []
        for formula in slide.get("formula", []):
            equation_markup = escape_text(formula.get("ppt", formula.get("eq", "")))
            if formula.get("tex"):
                equation_markup = (
                    f'<div class="equation math-display" data-math = "1">'
                    f'\\[{escape_text(formula["tex"])}\\]'
                    f'</div>'
                )
            else:
                equation_markup = f'<div class="equation">{equation_markup}</div>'
            formulas.append(
                f'<div class="formula-line">'
                f'{equation_markup}'
                f'<p>{escape_text(formula["desc"])}</p>'
                f'</div>'
            )
        definitions = "".join(
            (
                f'<span class="definition-pill" data-math = "1">'
                f'\\({escape_text(item["tex"])}\\): {escape_text(item["desc"])}'
                f'</span>'
            )
            for item in slide.get("definitions", [])
        )
        return (
            f'<div class="formula-panel">{"".join(formulas)}</div>'
            f'<div class="definition-row">{definitions}</div>'
        )
    if layout == "factors":
        cards = "".join(render_card(card) for card in slide.get("cards", []))
        return f'<div class="factor-grid">{cards}</div>'
    if layout == "joint_chart":
        return render_joint_svg()
    if layout == "dataset_matrix":
        return render_dataset_matrix(slide)
    if layout == "comparison":
        cards = "".join(render_card(card) for card in slide.get("cards", []))
        return f'<div class="compare-grid">{cards}</div>'
    if layout == "literature":
        cards = "".join(render_card(card) for card in slide.get("cards", []))
        return f'<div class="literature-grid">{cards}</div>'
    if layout == "takeaways":
        cards = "".join(render_card(card) for card in slide.get("cards", []))
        return f'<div class="takeaway-grid">{cards}</div>'
    return ""


def render_slide(slide, total):
    """Render one HTML slide section.

    Args:
        slide: Slide dictionary to render.
        total: Total slide count for footer numbering.

    Returns:
        str: HTML markup for one slide.
    """
    active = " is-active" if slide["id"] == 1 else ""
    if slide["layout"] == "cover":
        header = ""
    else:
        header = (
            f'<div class="slide-head">'
            f'<p class="kicker">{escape_text(slide["kicker"])}</p>'
            f'<span>{slide["id"]:02d} / {total:02d}</span>'
            f'</div>'
            f'<h2>{escape_text(slide["title"])}</h2>'
            f'<p class="subtitle">{escape_text(slide["subtitle"])}</p>'
        )
    return (
        f'<section class="slide{active}" data-title="{escape_text(slide["title"])}">'
        f'{header}'
        f'{render_visual(slide)}'
        f'{render_explanation(slide)}'
        f'<footer><span>{escape_text(slide["takeaway"])}</span>'
        f'<span class="slide-number" data-current="{slide["id"]}" data-total="{total}"></span></footer>'
        f'{render_notes(slide)}'
        f'</section>'
    )


def build_html():
    """Build the complete HTML deck.

    Returns:
        str: Complete HTML document.
    """
    total = len(SLIDES)
    slides = "\n".join(render_slide(slide, total) for slide in SLIDES)
    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>CS6493 Math Reasoning Presentation</title>
<link rel="stylesheet" href="style.css">
<script>
window.MathJax = {{
  tex: {{
    inlineMath: [['\\\\(', '\\\\)']],
    displayMath: [['\\\\[', '\\\\]']]
  }},
  svg: {{
    fontCache: 'global'
  }}
}};
</script>
<script defer src="https://cdn.jsdelivr.net/npm/mathjax@3/es5/tex-svg.js"></script>
</head>
<body>
<div class="deck">
{slides}
</div>
<div class="keyboard-hint">S presenter notes | N notes drawer | O overview | F fullscreen | left/right navigate</div>
<script src="runtime.js"></script>
</body>
</html>
"""


def build_css():
    """Build the CSS stylesheet for the HTML deck.

    Returns:
        str: Complete CSS stylesheet.
    """
    return """
:root {
  --paper: #f7f9fc;
  --white: #ffffff;
  --ink: #18212f;
  --muted: #647084;
  --line: #ced6e3;
  --blue: #1e5b9e;
  --teal: #0f766e;
  --amber: #b7791f;
  --red: #b42318;
  --green: #2f855a;
  --slate: #344256;
  --soft-blue: #eaf2ff;
  --soft-teal: #e7f6f4;
  --soft-amber: #fff3d6;
  --soft-red: #feeceb;
  --soft-green: #eaf7ef;
  --shadow: 0 18px 48px rgba(24, 33, 47, 0.10);
}

* {
  box-sizing: border-box;
}

html,
body {
  width: 100%;
  height: 100%;
  margin: 0;
  overflow: hidden;
  background: var(--paper);
  color: var(--ink);
  font-family: Arial, Helvetica, sans-serif;
}

.deck {
  position: relative;
  width: 100vw;
  height: 100vh;
}

.slide {
  position: absolute;
  inset: 0;
  display: none;
  padding: 58px 76px 52px;
  background:
    linear-gradient(90deg, rgba(30, 91, 158, 0.09) 1px, transparent 1px),
    linear-gradient(0deg, rgba(30, 91, 158, 0.07) 1px, transparent 1px),
    var(--paper);
  background-size: 48px 48px;
}

.slide.is-active {
  display: block;
}

.slide::before {
  content: "";
  position: absolute;
  top: 0;
  left: 0;
  width: 12px;
  height: 100%;
  background: linear-gradient(180deg, var(--blue), var(--teal), var(--amber));
}

.slide-head {
  display: flex;
  align-items: center;
  justify-content: space-between;
  margin-bottom: 14px;
}

.kicker {
  margin: 0;
  color: var(--blue);
  font-family: "Courier New", Courier, monospace;
  font-size: 13px;
  font-weight: 700;
  letter-spacing: 0.12em;
  text-transform: uppercase;
}

.slide-head span,
footer,
.keyboard-hint {
  color: var(--muted);
  font-family: "Courier New", Courier, monospace;
  font-size: 12px;
}

h1,
h2,
h3,
p {
  margin: 0;
}

h1 {
  max-width: 1040px;
  color: var(--ink);
  font-family: Georgia, "Times New Roman", serif;
  font-size: clamp(46px, 6vw, 78px);
  line-height: 1.05;
  letter-spacing: 0;
}

h2 {
  max-width: 1050px;
  color: var(--ink);
  font-family: Georgia, "Times New Roman", serif;
  font-size: clamp(34px, 4.2vw, 54px);
  line-height: 1.12;
  letter-spacing: 0;
}

.subtitle {
  max-width: 980px;
  margin-top: 14px;
  color: var(--muted);
  font-size: 19px;
  line-height: 1.45;
}

.cover-layout {
  display: grid;
  grid-template-columns: 1fr 260px;
  gap: 52px;
  align-items: center;
  height: calc(100vh - 120px);
}

.cover-mark {
  display: grid;
  place-items: center;
  width: 236px;
  height: 236px;
  border: 2px solid var(--blue);
  background: var(--white);
  box-shadow: var(--shadow);
}

.cover-mark span {
  color: var(--blue);
  font-family: Georgia, "Times New Roman", serif;
  font-size: 92px;
  font-weight: 700;
  line-height: 0.85;
}

.cover-mark small {
  color: var(--muted);
  font-family: "Courier New", Courier, monospace;
  font-size: 15px;
  text-transform: uppercase;
}

.chips {
  display: flex;
  flex-wrap: wrap;
  gap: 10px;
  margin-top: 26px;
}

.chip {
  display: inline-flex;
  align-items: center;
  min-height: 30px;
  padding: 6px 12px;
  border: 1px solid var(--line);
  background: var(--white);
  color: var(--slate);
  font-family: "Courier New", Courier, monospace;
  font-size: 12px;
  font-weight: 700;
}

.card-grid,
.setup-grid,
.lane-grid,
.factor-grid,
.literature-grid,
.takeaway-grid,
.compare-grid {
  display: grid;
  gap: 20px;
  margin-top: 42px;
}

.card-grid.count-3,
.card-grid.count-6,
.setup-grid,
.factor-grid,
.literature-grid,
.takeaway-grid {
  grid-template-columns: repeat(3, minmax(0, 1fr));
}

.card-grid.count-4 {
  grid-template-columns: repeat(4, minmax(0, 1fr));
}

.card,
.lane,
.step,
.loop-node {
  min-height: 178px;
  padding: 24px 24px 22px;
  border: 1px solid var(--line);
  border-top-width: 5px;
  background: var(--white);
  box-shadow: var(--shadow);
}

.card h3,
.lane h3,
.step h3,
.loop-node h3 {
  margin-bottom: 12px;
  color: var(--ink);
  font-size: 22px;
  line-height: 1.2;
}

.card p,
.lane p,
.step p,
.loop-node p {
  color: var(--muted);
  font-size: 16px;
  line-height: 1.45;
}

.mini-equation {
  display: block;
  margin-top: 8px;
  color: var(--blue);
  font-family: Georgia, "Times New Roman", serif;
  font-size: 14px;
  font-style: italic;
  font-weight: 700;
  line-height: 1.35;
}

.accent-blue {
  border-top-color: var(--blue);
}

.accent-teal {
  border-top-color: var(--teal);
}

.accent-amber {
  border-top-color: var(--amber);
}

.accent-red {
  border-top-color: var(--red);
}

.accent-green {
  border-top-color: var(--green);
}

.callout {
  margin-top: 26px;
  padding: 18px 22px;
  border-left: 6px solid var(--teal);
  background: var(--soft-teal);
  color: var(--slate);
  font-size: 19px;
  font-weight: 700;
}

.lane-grid {
  grid-template-columns: repeat(3, minmax(0, 1fr));
}

.lane .chips {
  margin-top: 0;
  margin-bottom: 18px;
}

.step-grid {
  display: grid;
  gap: 18px;
  margin-top: 42px;
}

.step-grid.count-3 {
  grid-template-columns: repeat(3, minmax(0, 1fr));
}

.step-grid.count-4 {
  grid-template-columns: repeat(4, minmax(0, 1fr));
}

.step span {
  display: inline-block;
  margin-bottom: 18px;
  color: var(--blue);
  font-family: Georgia, "Times New Roman", serif;
  font-size: 52px;
  font-weight: 700;
}

.loop-grid {
  display: grid;
  grid-template-columns: 1fr auto 1fr auto 1fr auto 1fr;
  gap: 12px;
  align-items: center;
  margin-top: 42px;
}

.loop-grid b {
  color: var(--muted);
  font-family: "Courier New", Courier, monospace;
  font-size: 18px;
}

.loop-node {
  min-height: 240px;
}

.loop-node span {
  display: inline-grid;
  place-items: center;
  width: 36px;
  height: 36px;
  margin-bottom: 20px;
  background: var(--soft-blue);
  color: var(--blue);
  font-family: "Courier New", Courier, monospace;
  font-weight: 700;
}

.formula-panel {
  display: grid;
  gap: 12px;
  margin-top: 30px;
  padding: 22px 26px;
  border: 1px solid var(--line);
  background: var(--white);
  box-shadow: var(--shadow);
}

.formula-line {
  padding: 14px 18px;
  border-left: 5px solid var(--blue);
  background: var(--paper);
}

.formula-line .equation {
  color: var(--ink);
  font-family: Georgia, "Times New Roman", serif;
  font-size: 26px;
  font-style: italic;
  font-weight: 700;
}

.math-display .mjx-container,
.math-inline .mjx-container,
.definition-pill .mjx-container {
  margin: 0 !important;
}

.math-inline .mjx-container {
  font-size: 0.98em !important;
}

.formula-line p {
  margin-top: 6px;
  color: var(--muted);
  font-size: 14px;
  line-height: 1.36;
}

.definition-row {
  display: flex;
  flex-wrap: wrap;
  gap: 8px;
  margin-top: 14px;
}

.definition-row span {
  padding: 6px 10px;
  border: 1px solid var(--line);
  background: var(--white);
  color: var(--slate);
  font-family: Georgia, "Times New Roman", serif;
  font-size: 13px;
  font-style: italic;
}

.dataset-matrix {
  display: grid;
  grid-template-columns: repeat(2, minmax(0, 1fr));
  gap: 20px;
  margin-top: 30px;
}

.matrix-panel {
  padding: 18px 18px 14px;
  border: 1px solid var(--line);
  border-top-width: 5px;
  background: var(--white);
  box-shadow: var(--shadow);
}

.matrix-panel h3 {
  margin-bottom: 12px;
  color: var(--ink);
  font-size: 20px;
}

.score-table {
  width: 100%;
  border-collapse: collapse;
  table-layout: fixed;
}

.score-table th,
.score-table td {
  padding: 8px 8px;
  border: 1px solid var(--line);
  text-align: center;
  font-size: 13px;
}

.score-table thead th {
  background: var(--paper);
  color: var(--slate);
  font-family: "Courier New", Courier, monospace;
  font-weight: 700;
}

.score-table tbody th {
  width: 28%;
  background: var(--white);
  color: var(--ink);
  font-size: 12px;
  text-align: left;
}

.score-cell {
  font-family: "Courier New", Courier, monospace;
  font-weight: 700;
}

.band-top {
  background: var(--soft-green);
  color: var(--green);
}

.band-high {
  background: var(--soft-teal);
  color: var(--teal);
}

.band-mid {
  background: var(--soft-blue);
  color: var(--blue);
}

.band-low {
  background: var(--soft-amber);
  color: var(--amber);
}

.band-risk {
  background: var(--soft-red);
  color: var(--red);
}

.explain-box {
  position: absolute;
  left: 76px;
  right: 76px;
  bottom: 70px;
  min-height: 58px;
  padding: 12px 16px 12px 18px;
  border: 1px solid var(--line);
  background: rgba(255, 255, 255, 0.88);
  color: var(--slate);
  box-shadow: 0 10px 26px rgba(24, 33, 47, 0.08);
}

.explain-box::before {
  content: "";
  position: absolute;
  left: 0;
  top: 0;
  bottom: 0;
  width: 5px;
  background: var(--blue);
}

.explain-box.accent-teal::before {
  background: var(--teal);
}

.explain-box.accent-amber::before {
  background: var(--amber);
}

.explain-box.accent-green::before {
  background: var(--green);
}

.explain-head {
  display: flex;
  align-items: center;
  justify-content: space-between;
  gap: 12px;
  margin-bottom: 6px;
}

.explain-tag {
  display: inline-flex;
  align-items: center;
  padding: 3px 9px;
  border-radius: 999px;
  background: var(--soft-blue);
  color: var(--blue);
  font-family: "Courier New", Courier, monospace;
  font-size: 11px;
  font-weight: 700;
  letter-spacing: 0.08em;
  text-transform: uppercase;
}

.explain-box.accent-teal .explain-tag {
  background: var(--soft-teal);
  color: var(--teal);
}

.explain-box.accent-amber .explain-tag {
  background: var(--soft-amber);
  color: var(--amber);
}

.explain-box.accent-green .explain-tag {
  background: var(--soft-green);
  color: var(--green);
}

.explain-meta {
  color: var(--muted);
  font-family: "Courier New", Courier, monospace;
  font-size: 11px;
  font-weight: 700;
  letter-spacing: 0.08em;
  text-transform: uppercase;
}

.explain-box p {
  margin: 0;
  color: var(--slate);
  font-size: 14px;
  line-height: 1.38;
}

.chart-panel {
  margin-top: 30px;
  padding: 22px 28px 18px;
  border: 1px solid var(--line);
  background: var(--white);
  box-shadow: var(--shadow);
}

.legend {
  color: var(--muted);
  font-size: 14px;
  font-weight: 700;
}

.dot {
  display: inline-block;
  width: 10px;
  height: 10px;
  margin: 0 6px 0 18px;
}

.dot.ds,
.bar.ds {
  fill: var(--blue);
  background: var(--blue);
}

.dot.qw,
.bar.qw {
  fill: var(--teal);
  background: var(--teal);
}

.joint-svg {
  width: 100%;
  height: min(50vh, 390px);
}

.axis {
  fill: var(--muted);
  font-family: "Courier New", Courier, monospace;
  font-size: 13px;
  font-weight: 700;
}

.value {
  fill: var(--ink);
  font-family: "Courier New", Courier, monospace;
  font-size: 13px;
  font-weight: 700;
}

.gridline {
  stroke: var(--line);
  stroke-width: 1;
}

.compare-grid {
  grid-template-columns: repeat(2, minmax(0, 1fr));
}

footer {
  position: absolute;
  left: 76px;
  right: 76px;
  bottom: 28px;
  display: flex;
  align-items: center;
  justify-content: space-between;
  gap: 18px;
}

footer span:first-child {
  max-width: 920px;
  color: var(--slate);
  font-family: Arial, Helvetica, sans-serif;
  font-size: 15px;
  font-weight: 700;
}

.notes {
  display: none;
}

.keyboard-hint {
  position: fixed;
  right: 16px;
  bottom: 8px;
  z-index: 50;
  pointer-events: none;
  opacity: 0.62;
}

.progress-bar {
  position: fixed;
  left: 0;
  right: 0;
  bottom: 0;
  z-index: 80;
  height: 5px;
  background: rgba(24, 33, 47, 0.12);
}

.progress-bar span {
  display: block;
  width: 0;
  height: 100%;
  background: linear-gradient(90deg, var(--blue), var(--teal), var(--amber));
}

.notes-overlay {
  position: fixed;
  left: 72px;
  right: 72px;
  bottom: 44px;
  z-index: 120;
  display: none;
  max-height: 42vh;
  overflow: auto;
  padding: 22px 26px;
  border: 1px solid var(--line);
  background: rgba(255, 255, 255, 0.96);
  box-shadow: var(--shadow);
  color: var(--ink);
  font-size: 18px;
  line-height: 1.5;
}

.notes-overlay.open {
  display: block;
}

.notes-overlay p + p {
  margin-top: 14px;
}

.overview {
  position: fixed;
  inset: 0;
  z-index: 160;
  display: none;
  grid-template-columns: repeat(3, 1fr);
  gap: 18px;
  padding: 34px;
  overflow: auto;
  background: rgba(24, 33, 47, 0.92);
}

.overview.open {
  display: grid;
}

.overview .thumb {
  cursor: pointer;
  border: 1px solid rgba(255, 255, 255, 0.22);
  background: var(--white);
}

@media (max-width: 900px) {
  .slide {
    padding: 40px 34px 56px;
  }

  .cover-layout,
  .card-grid.count-3,
  .card-grid.count-4,
  .card-grid.count-6,
  .setup-grid,
  .lane-grid,
  .factor-grid,
  .dataset-matrix,
  .literature-grid,
  .takeaway-grid,
  .compare-grid,
  .step-grid.count-3,
  .step-grid.count-4,
  .loop-grid {
    grid-template-columns: 1fr;
  }

  .cover-mark {
    display: none;
  }

  footer {
    left: 34px;
    right: 34px;
  }

  .explain-box {
    left: 34px;
    right: 34px;
    font-size: 13px;
  }
}
"""


def build_manifest():
    """Build the slide manifest data.

    Returns:
        list: JSON-serializable slide manifest records.
    """
    return [
        {
            "slide": slide["id"],
            "title": slide["title"],
            "kicker": slide["kicker"],
            "layout": slide["layout"],
            "role": "main",
            "takeaway": slide["takeaway"],
            "explanation": slide.get("explanation", ""),
            "speaker_notes": slide.get("notes", []),
            "has_speaker_notes": bool(slide.get("notes")),
        }
        for slide in SLIDES
    ]


def build_speaker_notes():
    """Build a standalone speaker-notes markdown file.

    Returns:
        str: Markdown speaker notes grouped by slide.
    """
    sections = ["# CS6493 Speaker Notes", ""]
    for slide in SLIDES:
        sections.append(f"## {slide['id']:02d}. {slide['title']}")
        sections.append(f"- Kicker: {slide['kicker']}")
        sections.append(f"- Takeaway: {slide['takeaway']}")
        for note in slide.get("notes", []):
            sections.append(f"- {note}")
        sections.append("")
    return "\n".join(sections)


def build_readme():
    """Build the deck README.

    Returns:
        str: Markdown documentation for the generated deck.
    """
    return f"""# CS6493 Math Reasoning Presentation Deck

This folder contains an English classroom presentation based on `delivery/reports/final_report_en.tex`.

## Files

- `index.html`: keyboard-driven HTML deck with hidden speaker notes.
- `style.css`: Academic Tech styling for the HTML deck.
- `runtime.js`: local copy of the html-ppt runtime for navigation, overview, notes, and presenter mode.
- `{PPTX_NAME}`: editable PowerPoint deck generated with `python-pptx`.
- `slide_manifest.json`: slide titles, layouts, and takeaways.
- `speaker_notes.md`: standalone notes for rehearsal and presenter reference.
- `build_deck.py`: single-source generator for the HTML, manifest, README, and PPTX.
- `validate_deck.py`: validation helper for PPTX structure and 16:9 HTML screenshots.

## Usage

Open `index.html` in a browser for presentation. Keyboard controls:

- `left` / `right` / `space`: navigate slides
- `S`: open presenter mode
- `N`: open the notes drawer
- `O`: open slide overview
- `F`: fullscreen

Regenerate all files:

```bash
python delivery/ppt_slides/build_deck.py
```

Regenerate only HTML assets:

```bash
python delivery/ppt_slides/build_deck.py --skip-pptx
```

Validate the generated deck:

```bash
python delivery/ppt_slides/validate_deck.py
```

The PowerPoint file is intentionally built with editable text boxes, shapes, tables, and charts rather than full-slide screenshots.
"""


def add_shape(slide, shape_type, x, y, width, height, fill, line = None):
    """Add a colored shape to a slide.

    Args:
        slide: PowerPoint slide object.
        shape_type: MSO_SHAPE shape type.
        x: Left position in inches.
        y: Top position in inches.
        width: Shape width in inches.
        height: Shape height in inches.
        fill: Fill color palette key.
        line: Optional line color palette key.

    Returns:
        pptx.shapes.autoshape.Shape: Created PowerPoint shape.
    """
    shape = slide.shapes.add_shape(
        shape_type,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color(fill))
    if line:
        shape.line.color.rgb = rgb(color(line))
    else:
        shape.line.fill.background()
    return shape


def add_text(slide, text, x, y, width, height, size = 18, bold = False, fill = "ink", align = PP_ALIGN.LEFT, font = "Arial"):
    """Add text to a PowerPoint slide.

    Args:
        slide: PowerPoint slide object.
        text: Text content.
        x: Left position in inches.
        y: Top position in inches.
        width: Text box width in inches.
        height: Text box height in inches.
        size: Font size in points.
        bold: Whether the font should be bold.
        fill: Text color palette key.
        align: Paragraph alignment.
        font: Font family name.

    Returns:
        pptx.shapes.autoshape.Shape: Created text box shape.
    """
    shape = slide.shapes.add_textbox(
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height)
    )
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    lines = str(text).split("\n")
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color(fill))
    return shape


def add_header(slide, slide_data, total):
    """Add the repeated slide header.

    Args:
        slide: PowerPoint slide object.
        slide_data: Slide dictionary for the current slide.
        total: Total number of slides.
    """
    add_text(
        slide,
        slide_data["kicker"].upper(),
        0.72,
        0.38,
        5.8,
        0.22,
        size = 8,
        bold = True,
        fill = "blue",
        font = "Courier New"
    )
    add_text(
        slide,
        f'{slide_data["id"]:02d} / {total:02d}',
        11.7,
        0.38,
        0.9,
        0.22,
        size = 8,
        bold = True,
        fill = "muted",
        align = PP_ALIGN.RIGHT,
        font = "Courier New"
    )
    add_text(
        slide,
        slide_data["title"],
        0.72,
        0.72,
        11.8,
        0.82,
        size = 25,
        bold = True,
        fill = "ink",
        font = "Georgia"
    )
    add_text(
        slide,
        slide_data["subtitle"],
        0.74,
        1.52,
        11.4,
        0.38,
        size = 12,
        fill = "muted"
    )


def add_footer(slide, slide_data):
    """Add the repeated takeaway footer.

    Args:
        slide: PowerPoint slide object.
        slide_data: Slide dictionary for the current slide.
    """
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0.72, 6.95, 11.6, 0.01, "line")
    add_text(
        slide,
        slide_data["takeaway"],
        0.74,
        7.02,
        10.8,
        0.26,
        size = 8,
        bold = True,
        fill = "slate"
    )


def add_explanation_box(slide, slide_data):
    """Add the visible academic explanation box to a PowerPoint slide.

    Args:
        slide: PowerPoint slide object.
        slide_data: Slide dictionary for the current slide.
    """
    explanation = slide_data.get("explanation", "")
    if not explanation:
        return
    label, meta, accent = get_explanation_meta(slide_data)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0.76, 6.18, 11.78, 0.6, "white", "line")
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0.76, 6.18, 0.06, 0.6, accent)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.92, 6.27, 1.45, 0.16, f"soft_{accent}" if accent in {"blue", "teal", "amber", "green"} else "soft_blue")
    add_text(
        slide,
        label.upper(),
        0.98,
        6.305,
        1.28,
        0.08,
        size = 6,
        bold = True,
        fill = accent,
        font = "Courier New",
        align = PP_ALIGN.CENTER
    )
    add_text(
        slide,
        meta,
        10.92,
        6.305,
        1.2,
        0.08,
        size = 6,
        bold = True,
        fill = "muted",
        font = "Courier New",
        align = PP_ALIGN.RIGHT
    )
    add_text(
        slide,
        explanation,
        0.92,
        6.48,
        11.26,
        0.16,
        size = 7,
        bold = False,
        fill = "slate"
    )


def score_fill(score):
    """Map a joint-score value to slide fill and font colors.

    Args:
        score: Joint-score value between zero and one.

    Returns:
        tuple: (fill color name, text color name)
    """
    band = score_band(score)
    if band == "top":
        return "soft_green", "green"
    if band == "high":
        return "soft_teal", "teal"
    if band == "mid":
        return "soft_blue", "blue"
    if band == "low":
        return "soft_amber", "amber"
    return "soft_red", "red"


def add_card(slide, card, x, y, width, height):
    """Add an editable card to a PowerPoint slide.

    Args:
        slide: PowerPoint slide object.
        card: Card dictionary with title, body, and color.
        x: Left position in inches.
        y: Top position in inches.
        width: Card width in inches.
        height: Card height in inches.
    """
    card_color = card.get("color", "blue")
    title = card.get("title", card.get("name", ""))
    body = card.get("body", card.get("short", ""))
    formula = card.get("formula")
    formula_ppt = card.get("formula_ppt")
    if formula_ppt:
        body = f"{body}\n{formula_ppt}"
    elif formula:
        body = f"{body}\n{standard_formula(formula)}"
    add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, width, height, "white", "line")
    add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, width, 0.07, card_color)
    add_text(
        slide,
        title,
        x + 0.18,
        y + 0.22,
        width - 0.36,
        0.34,
        size = 15,
        bold = True,
        fill = "ink"
    )
    add_text(
        slide,
        body,
        x + 0.18,
        y + 0.72,
        width - 0.36,
        height - 0.84,
        size = 10,
        fill = "muted"
    )


def add_cards_layout(slide, cards, x, y, width, height):
    """Add a responsive card grid to a PowerPoint slide.

    Args:
        slide: PowerPoint slide object.
        cards: List of card dictionaries.
        x: Left position in inches.
        y: Top position in inches.
        width: Total grid width in inches.
        height: Total grid height in inches.
    """
    count = len(cards)
    if count == 6:
        cols = 3
    elif count == 4:
        cols = 4
    elif count == 2:
        cols = 2
    else:
        cols = 3
    gap = 0.18
    card_width = (width - gap * (cols - 1)) / cols
    card_height = height
    for index, card in enumerate(cards):
        col = index % cols
        row = index // cols
        y_offset = y + row * (card_height + 0.22)
        add_card(slide, card, x + col * (card_width + gap), y_offset, card_width, card_height)


def add_cover(slide, slide_data):
    """Add the cover slide layout.

    Args:
        slide: PowerPoint slide object.
        slide_data: Slide dictionary for the cover.
    """
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0.0, 0.0, 0.16, 7.5, "blue")
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0.16, 0.0, 0.06, 7.5, "teal")
    add_text(
        slide,
        slide_data["kicker"].upper(),
        0.78,
        0.76,
        5.0,
        0.28,
        size = 9,
        bold = True,
        fill = "blue",
        font = "Courier New"
    )
    add_text(
        slide,
        slide_data["title"],
        0.74,
        1.2,
        8.6,
        1.9,
        size = 34,
        bold = True,
        fill = "ink",
        font = "Georgia"
    )
    add_text(
        slide,
        slide_data["subtitle"],
        0.78,
        3.3,
        7.6,
        0.48,
        size = 14,
        fill = "muted"
    )
    for index, tag in enumerate(slide_data["tags"]):
        add_shape(slide, MSO_SHAPE.RECTANGLE, 0.78, 4.16 + index * 0.46, 4.9, 0.3, "white", "line")
        add_text(
            slide,
            tag,
            0.92,
            4.22 + index * 0.46,
            4.6,
            0.18,
            size = 8,
            bold = True,
            fill = "slate",
            font = "Courier New"
        )
    add_shape(slide, MSO_SHAPE.RECTANGLE, 10.45, 1.54, 1.8, 1.8, "white", "blue")
    add_text(
        slide,
        str(len(SLIDES)),
        10.76,
        1.84,
        1.18,
        0.72,
        size = 42,
        bold = True,
        fill = "blue",
        align = PP_ALIGN.CENTER,
        font = "Georgia"
    )
    add_text(
        slide,
        "SLIDES",
        10.86,
        2.58,
        0.96,
        0.18,
        size = 8,
        bold = True,
        fill = "muted",
        align = PP_ALIGN.CENTER,
        font = "Courier New"
    )
    add_footer(slide, slide_data)


def add_setup_layout(slide, slide_data):
    """Add the experimental setup layout.

    Args:
        slide: PowerPoint slide object.
        slide_data: Slide dictionary with setup cards.
    """
    add_cards_layout(slide, slide_data["cards"], 0.82, 2.22, 11.5, 2.15)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 1.8, 4.92, 9.3, 0.58, "soft_teal", "teal")
    add_text(
        slide,
        "Total per model-method run: 130 fixed questions",
        2.06,
        5.08,
        8.8,
        0.24,
        size = 16,
        bold = True,
        fill = "teal",
        align = PP_ALIGN.CENTER
    )


def add_taxonomy_layout(slide, slide_data):
    """Add the method taxonomy layout.

    Args:
        slide: PowerPoint slide object.
        slide_data: Slide dictionary with lane data.
    """
    gap = 0.24
    lane_width = (11.5 - gap * 2) / 3
    for index, lane in enumerate(slide_data["lanes"]):
        x = 0.82 + index * (lane_width + gap)
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, 2.12, lane_width, 3.58, "white", "line")
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, 2.12, lane_width, 0.08, lane["color"])
        add_text(
            slide,
            lane["title"],
            x + 0.18,
            2.36,
            lane_width - 0.36,
            0.32,
            size = 15,
            bold = True,
            fill = "ink"
        )
        add_text(
            slide,
            "\n".join(lane["methods"]),
            x + 0.18,
            2.94,
            lane_width - 0.36,
            0.84,
            size = 11,
            bold = True,
            fill = lane["color"],
            font = "Courier New"
        )
        add_text(
            slide,
            lane["body"],
            x + 0.18,
            4.08,
            lane_width - 0.36,
            1.1,
            size = 10,
            fill = "muted"
        )


def add_process_layout(slide, slide_data):
    """Add a process or roadmap layout.

    Args:
        slide: PowerPoint slide object.
        slide_data: Slide dictionary with steps.
    """
    steps = slide_data["steps"]
    count = len(steps)
    gap = 0.18
    card_width = (11.5 - gap * (count - 1)) / count
    for index, step in enumerate(steps):
        x = 0.82 + index * (card_width + gap)
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, 2.2, card_width, 3.1, "white", "line")
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, 2.2, card_width, 0.08, step["color"])
        add_text(
            slide,
            f"{index + 1:02d}",
            x + 0.18,
            2.52,
            card_width - 0.36,
            0.56,
            size = 26,
            bold = True,
            fill = step["color"],
            font = "Georgia"
        )
        add_text(
            slide,
            step["title"],
            x + 0.18,
            3.28,
            card_width - 0.36,
            0.34,
            size = 14,
            bold = True,
            fill = "ink"
        )
        add_text(
            slide,
            step["body"],
            x + 0.18,
            3.84,
            card_width - 0.36,
            1.02,
            size = 9,
            fill = "muted"
        )


def add_formula_layout(slide, slide_data):
    """Add the score formula layout.

    Args:
        slide: PowerPoint slide object.
        slide_data: Slide dictionary with formula strings.
    """
    add_shape(slide, MSO_SHAPE.RECTANGLE, 1.0, 2.04, 11.0, 3.42, "white", "line")
    for index, formula in enumerate(slide_data["formula"]):
        y = 2.26 + index * 0.92
        add_shape(slide, MSO_SHAPE.RECTANGLE, 1.35, y, 10.3, 0.72, "paper", "line")
        add_shape(slide, MSO_SHAPE.RECTANGLE, 1.35, y, 0.06, 0.72, "blue")
        add_text(
            slide,
            formula["ppt"],
            1.58,
            y + 0.11,
            9.86,
            0.24,
            size = 17,
            bold = True,
            fill = "ink",
            font = "Georgia"
        )
        add_text(
            slide,
            formula["desc"],
            1.58,
            y + 0.43,
            9.86,
            0.2,
            size = 8,
            fill = "muted"
        )
    gap = 0.14
    card_width = (11.0 - gap * 2) / 3
    for index, item in enumerate(slide_data.get("definitions", [])):
        row = index // 3
        col = index % 3
        x = 1.0 + col * (card_width + gap)
        y = 5.58 + row * 0.28
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, card_width, 0.2, "soft_teal", "teal")
        add_text(
            slide,
            f'{item["ppt"]}: {item["desc"]}',
            x + 0.08,
            y + 0.045,
            card_width - 0.16,
            0.1,
            size = 6.5,
            bold = True,
            fill = "teal",
            font = "Georgia",
            align = PP_ALIGN.CENTER
        )


def add_joint_chart_layout(slide):
    """Add the editable joint score chart layout.

    Args:
        slide: PowerPoint slide object.
    """
    chart_data = CategoryChartData()
    chart_data.categories = METHODS
    chart_data.add_series("DeepSeek joint score", JOINT_DATA["DeepSeek"])
    chart_data.add_series("Qwen joint score", JOINT_DATA["Qwen"])
    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(0.9),
        Inches(2.05),
        Inches(10.95),
        Inches(3.82),
        chart_data
    )
    chart = chart_shape.chart
    chart.has_legend = True
    chart.legend.include_in_layout = False
    chart.value_axis.minimum_scale = 0.0
    chart.value_axis.maximum_scale = 0.75
    chart.value_axis.tick_labels.font.size = Pt(8)
    chart.category_axis.tick_labels.font.size = Pt(9)
    chart.plots[0].has_data_labels = True
    chart.plots[0].data_labels.number_format = "0.000"
    chart.plots[0].data_labels.font.size = Pt(8)
    for series_index, series in enumerate(chart.series):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = rgb(color("blue" if series_index == 0 else "teal"))


def add_dataset_matrix_layout(slide, slide_data):
    """Add side-by-side dataset joint-score tables.

    Args:
        slide: PowerPoint slide object.
        slide_data: Slide dictionary containing dataset matrix panels.
    """
    headers = [
        "Method",
        DATASET_LABELS["math500"],
        DATASET_LABELS["gsm8k"],
        DATASET_LABELS["aime2024"],
    ]
    table_width = 5.48
    positions = [0.82, 6.03]
    for panel_index, panel in enumerate(slide_data.get("panels", [])):
        x = positions[panel_index]
        accent = panel["accent"]
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, 2.08, table_width, 3.78, "white", "line")
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, 2.08, table_width, 0.08, accent)
        add_text(
            slide,
            panel["title"],
            x + 0.16,
            2.22,
            table_width - 0.32,
            0.22,
            size = 12,
            bold = True,
            fill = "ink"
        )
        shape = slide.shapes.add_table(
            1 + len(panel["rows"]),
            len(headers),
            Inches(x + 0.12),
            Inches(2.56),
            Inches(table_width - 0.24),
            Inches(2.96)
        )
        table = shape.table
        for col, header in enumerate(headers):
            cell = table.cell(0, col)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(color("paper"))
            cell.text = header
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Courier New"
                    run.font.size = Pt(7.5)
                    run.font.bold = True
                    run.font.color.rgb = rgb(color("slate"))
        for row_index, row in enumerate(panel["rows"], start = 1):
            method_cell = table.cell(row_index, 0)
            method_cell.fill.solid()
            method_cell.fill.fore_color.rgb = rgb(color("white"))
            method_cell.text = row["method"]
            for paragraph in method_cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(7.5)
                    run.font.bold = True
                    run.font.color.rgb = rgb(color("ink"))
            for score_index, score in enumerate(row["scores"], start = 1):
                cell = table.cell(row_index, score_index)
                fill_name, text_name = score_fill(score)
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(color(fill_name))
                cell.text = f"{score:.3f}"
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Courier New"
                        run.font.size = Pt(7.5)
                        run.font.bold = True
                        run.font.color.rgb = rgb(color(text_name))


def add_literature_table(slide):
    """Add an editable literature connection table.

    Args:
        slide: PowerPoint slide object.
    """
    rows = [
        [
            "Prompting lineage",
            "Wei et al. 2022; Wang et al. 2022, 2023; Madaan et al. 2024",
            "Reasoning elicitation, multi-path voting, critique, and explicit planning",
        ],
        [
            "Tool-grounded solving",
            "Gou et al. 2023 (ToRA)",
            "Reasoning can call reliable computation, observe it, and continue",
        ],
        [
            "Efficiency-aware diagnosis",
            "Wang et al. 2025; Yang et al. 2025",
            "Reflection-token cost, overlong reasoning, and answer-aware stopping",
        ],
    ]
    shape = slide.shapes.add_table(
        4,
        3,
        Inches(0.86),
        Inches(2.18),
        Inches(11.42),
        Inches(3.24)
    )
    table = shape.table
    headers = [
        "Thread",
        "References",
        "How it informs this project",
    ]
    for col, header in enumerate(headers):
        cell = table.cell(0, col)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(color("soft_blue"))
        cell.text = header
        for paragraph in cell.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(10)
                run.font.bold = True
                run.font.color.rgb = rgb(color("blue"))
    for row_index, row in enumerate(rows, start = 1):
        for col, value in enumerate(row):
            cell = table.cell(row_index, col)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(color("white"))
            cell.text = value
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9)
                    run.font.color.rgb = rgb(color("ink" if col == 0 else "muted"))


def add_slide_body(slide, slide_data):
    """Add layout-specific body content to a slide.

    Args:
        slide: PowerPoint slide object.
        slide_data: Slide dictionary for the current slide.
    """
    layout = slide_data["layout"]
    if layout == "cover":
        add_cover(slide, slide_data)
    elif layout == "setup":
        add_setup_layout(slide, slide_data)
    elif layout == "taxonomy":
        add_taxonomy_layout(slide, slide_data)
    elif layout in {"process", "roadmap", "loop"}:
        add_process_layout(slide, slide_data)
    elif layout == "formula":
        add_formula_layout(slide, slide_data)
    elif layout == "joint_chart":
        add_joint_chart_layout(slide)
    elif layout == "dataset_matrix":
        add_dataset_matrix_layout(slide, slide_data)
    elif layout == "literature":
        add_literature_table(slide)
    elif layout in {"cards", "method_cards", "dataset_cards", "comparison", "takeaways", "factors"}:
        card_count = len(slide_data.get("cards", []))
        if card_count == 6:
            height = 1.42
        elif card_count < 4:
            height = 2.25
        else:
            height = 2.55
        add_cards_layout(slide, slide_data["cards"], 0.82, 2.34, 11.5, height)


def build_pptx(path):
    """Build the editable PowerPoint deck.

    Args:
        path: Destination path for the generated PPTX file.
    """
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    blank_layout = presentation.slide_layouts[6]
    total = len(SLIDES)
    for slide_data in SLIDES:
        slide = presentation.slides.add_slide(blank_layout)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(color("paper"))
        if slide_data["layout"] != "cover":
            add_shape(slide, MSO_SHAPE.RECTANGLE, 0.0, 0.0, 0.13, 7.5, "blue")
            add_shape(slide, MSO_SHAPE.RECTANGLE, 0.13, 0.0, 0.05, 7.5, "teal")
            add_header(slide, slide_data, total)
        add_slide_body(slide, slide_data)
        add_explanation_box(slide, slide_data)
        if slide_data["layout"] != "cover":
            add_footer(slide, slide_data)
    presentation.save(path)


def write_outputs(out_dir, skip_pptx = False):
    """Write all generated deck files.

    Args:
        out_dir: Output directory path.
        skip_pptx: Whether to skip PowerPoint generation.
    """
    out_dir.mkdir(parents = True, exist_ok = True)
    write_text(out_dir / "index.html", build_html())
    write_text(out_dir / "style.css", build_css())
    write_text(
        out_dir / "slide_manifest.json",
        json.dumps(build_manifest(), indent = 2, ensure_ascii = False) + "\n"
    )
    write_text(out_dir / "speaker_notes.md", build_speaker_notes())
    write_text(out_dir / "README.md", build_readme())
    runtime_source = Path("/Users/brench/.agents/skills/html-ppt/assets/runtime.js")
    runtime_target = out_dir / "runtime.js"
    runtime_target.write_text(runtime_source.read_text(encoding = "utf-8"), encoding = "utf-8")
    if not skip_pptx:
        build_pptx(out_dir / PPTX_NAME)
    logger.info("Wrote deck outputs to %s", out_dir)


def main():
    """Run the deck generator."""
    args = parse_args()
    write_outputs(args.out_dir, skip_pptx = args.skip_pptx)


if __name__ == "__main__":
    logging.basicConfig(
        level = logging.INFO,
        format = '%(asctime)s - %(name)s - %(levelname)s - %(message)s',
        handlers = [logging.StreamHandler()]
    )
    main()
